#!/usr/bin/env python3
"""Genera los documentos de ejemplo (ficticios, sin datos personales reales)
y la plantilla Word de muestra que se incluyen en la carpeta portátil.

Ejecutar una sola vez desde la raíz del repositorio:
    python scripts/generar_ejemplos.py
"""
from __future__ import annotations

import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from PIL import Image, ImageDraw, ImageFont
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

RUTA_FUENTE = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"

TEXTO_ACTA_1 = (
    "NOTARIA PUBLICA NUMERO 1 (FICTICIA)\n\n"
    "ESCRITURA NUMERO 99999\n\n"
    "En la Ciudad Ficticia, Estado de Jalisco, siendo el dia 10 de enero\n"
    "de 2023, ante mi, Licenciado Roberto Ejemplo Prueba, Notario Publico\n"
    "numero 1 de este Distrito, comparecen los senores Ana Ejemplo Muestra\n"
    "y Luis Prueba Demo, mayores de edad, mexicanos, con domicilio en\n"
    "Avenida Ficticia numero 200, Colonia Ejemplo, CP 45000, quienes\n"
    "otorgan el presente contrato de compraventa por la cantidad de\n"
    "$300,000.00 (trescientos mil pesos 00/100 M.N.), respecto de un\n"
    "inmueble con superficie de 250.00 m2, ubicado en Calle Muestra\n"
    "numero 300, con clave catastral 01-234-567-890, inscrito en el\n"
    "Registro Publico de la Propiedad bajo el Folio Real 123456.\n\n"
    "El inmueble colinda al norte con Juan Colindante Uno en 20.00 metros,\n"
    "al sur con Pedro Colindante Dos en 20.00 metros, al oriente con la\n"
    "Avenida Ficticia en 12.50 metros y al poniente con Maria Colindante\n"
    "Tres en 12.50 metros.\n\n"
    "Leido que les fue el presente instrumento, manifestaron su\n"
    "conformidad y lo firmaron ante mi, que doy fe.\n"
)

TEXTO_ACTA_2 = (
    "ACTA DE EJEMPLO NUMERO 2 (DATOS FICTICIOS)\n\n"
    "Expediente numero EJ-2024-0001, Volumen 15, Foja 320.\n\n"
    "Correo de contacto: contacto.ejemplo@notaria-ficticia.mx\n"
    "Porcentaje de copropiedad: 50%.\n"
)


def _crear_imagen(texto: str, ruta: Path, ancho=1275, alto=1650, tamano=20, angulo=0.0) -> Path:
    imagen = Image.new("RGB", (ancho, alto), "white")
    dibujo = ImageDraw.Draw(imagen)
    try:
        fuente = ImageFont.truetype(RUTA_FUENTE, tamano)
    except OSError:
        fuente = ImageFont.load_default()
    y = 100
    for linea in texto.split("\n"):
        dibujo.text((100, y), linea, fill="black", font=fuente)
        y += tamano + 14
    if angulo:
        imagen = imagen.rotate(angulo, fillcolor="white", expand=False)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    if ruta.suffix.lower() in (".jpg", ".jpeg"):
        imagen.save(ruta, dpi=(300, 300), quality=80)
    else:
        imagen.save(ruta, dpi=(300, 300))
    return ruta


def _crear_pdf_escaneado(rutas_imagenes: list[Path], ruta_salida: Path) -> Path:
    import fitz
    doc = fitz.open()
    for ruta_img in rutas_imagenes:
        img = Image.open(ruta_img)
        ancho_pt, alto_pt = img.width * 72 / 300, img.height * 72 / 300
        pagina = doc.new_page(width=ancho_pt, height=alto_pt)
        pagina.insert_image(pagina.rect, filename=str(ruta_img))
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ruta_salida))
    doc.close()
    return ruta_salida


def generar_ejemplos() -> None:
    carpeta_ejemplos = RAIZ / "ejemplos"
    carpeta_temp = RAIZ / "scripts" / "_tmp_ejemplos"
    carpeta_temp.mkdir(parents=True, exist_ok=True)

    img1 = _crear_imagen(TEXTO_ACTA_1, carpeta_temp / "pagina1.jpg")
    img2 = _crear_imagen(TEXTO_ACTA_2, carpeta_temp / "pagina2.jpg")
    _crear_pdf_escaneado([img1, img2], carpeta_ejemplos / "acta_ficticia_escaneada.pdf")

    _crear_imagen(TEXTO_ACTA_1, carpeta_ejemplos / "imagen_ficticia.png")
    _crear_imagen(TEXTO_ACTA_2, carpeta_ejemplos / "acta_ficticia_inclinada.png", angulo=4.5)

    import shutil
    shutil.rmtree(carpeta_temp, ignore_errors=True)
    print(f"Ejemplos generados en: {carpeta_ejemplos}")


def generar_plantilla_word() -> None:
    ruta = RAIZ / "resources" / "plantillas" / "Plantilla_Notarial.docx"
    ruta.parent.mkdir(parents=True, exist_ok=True)

    documento = docx.Document()
    for seccion in documento.sections:
        seccion.top_margin = Cm(2.5)
        seccion.bottom_margin = Cm(2.5)
        seccion.left_margin = Cm(3.0)
        seccion.right_margin = Cm(2.0)

        encabezado = seccion.header
        p_enc = encabezado.paragraphs[0]
        p_enc.text = "NOTARÍA PÚBLICA (nombre de ejemplo) — Sustituir por los datos reales de la notaría"
        p_enc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p_enc.runs:
            run.bold = True
            run.font.size = Pt(10)

        pie = seccion.footer
        p_pie = pie.paragraphs[0]
        p_pie.text = "Documento generado a partir de transcripción notarial — uso interno"
        p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p_pie.runs:
            run.font.size = Pt(8)
            run.italic = True

    titulo = documento.add_paragraph("TRANSCRIPCIÓN NOTARIAL")
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in titulo.runs:
        run.bold = True
        run.font.size = Pt(14)

    documento.add_paragraph(
        "A continuación se inserta la transcripción revisada del acta o "
        "documento notarial correspondiente."
    )
    documento.add_paragraph("")

    documento.save(str(ruta))
    print(f"Plantilla de ejemplo generada en: {ruta}")


if __name__ == "__main__":
    generar_ejemplos()
    generar_plantilla_word()
