from pathlib import Path

import docx
import pytest

from transcriptor_notarial.config import ConfiguracionApp
from transcriptor_notarial.core import docx_export


def test_genera_documento_word_legible(tmp_path: Path):
    config = ConfiguracionApp()
    paginas = [
        docx_export.PaginaTranscrita(numero=1, texto="Primer parrafo de prueba.\n\nSegundo parrafo."),
        docx_export.PaginaTranscrita(numero=2, texto="Contenido de la segunda pagina."),
    ]
    ruta_salida = tmp_path / "salida.docx"

    resultado = docx_export.generar_documento(paginas, config, ruta_salida)

    assert resultado.exists()
    documento = docx.Document(str(resultado))
    texto_completo = "\n".join(p.text for p in documento.paragraphs)
    assert "Primer parrafo de prueba." in texto_completo
    assert "Contenido de la segunda pagina." in texto_completo


def test_respeta_configuracion_de_fuente_y_tamano(tmp_path: Path):
    config = ConfiguracionApp(fuente_word="Arial", tamano_letra=14)
    paginas = [docx_export.PaginaTranscrita(numero=1, texto="Texto de prueba.")]
    ruta_salida = tmp_path / "salida.docx"

    docx_export.generar_documento(paginas, config, ruta_salida)

    documento = docx.Document(str(ruta_salida))
    parrafo = documento.paragraphs[0]
    assert parrafo.runs[0].font.name == "Arial"
    assert parrafo.runs[0].font.size.pt == 14


def test_salto_de_pagina_entre_paginas_del_original(tmp_path: Path):
    from docx.enum.text import WD_BREAK

    config = ConfiguracionApp()
    paginas = [
        docx_export.PaginaTranscrita(numero=1, texto="Pagina uno."),
        docx_export.PaginaTranscrita(numero=2, texto="Pagina dos."),
    ]
    ruta_salida = tmp_path / "salida.docx"
    docx_export.generar_documento(paginas, config, ruta_salida)

    documento = docx.Document(str(ruta_salida))
    xml = documento.element.xml
    assert "w:br" in xml  # existe al menos un salto (de página)


def test_modalidad_formato_semejante_genera_tabla(tmp_path: Path):
    config = ConfiguracionApp()
    paginas = [
        docx_export.PaginaTranscrita(
            numero=1, texto="", es_tabla=True,
            filas_tabla=[["Concepto", "Monto"], ["Terreno", "$100,000.00"]],
        )
    ]
    ruta_salida = tmp_path / "salida.docx"
    docx_export.generar_documento(
        paginas, config, ruta_salida, modalidad=docx_export.ModalidadExportacion.FORMATO_SEMEJANTE
    )

    documento = docx.Document(str(ruta_salida))
    assert len(documento.tables) == 1
    tabla = documento.tables[0]
    assert tabla.cell(0, 0).text == "Concepto"
    assert tabla.cell(1, 1).text == "$100,000.00"


def test_usa_plantilla_cuando_se_proporciona(tmp_path: Path):
    ruta_plantilla = tmp_path / "plantilla.docx"
    plantilla = docx.Document()
    plantilla.add_paragraph("ENCABEZADO DE LA NOTARIA")
    plantilla.save(str(ruta_plantilla))

    config = ConfiguracionApp()
    paginas = [docx_export.PaginaTranscrita(numero=1, texto="Cuerpo de la transcripcion.")]
    ruta_salida = tmp_path / "salida.docx"

    docx_export.generar_documento(paginas, config, ruta_salida, ruta_plantilla=ruta_plantilla)

    documento = docx.Document(str(ruta_salida))
    texto_completo = "\n".join(p.text for p in documento.paragraphs)
    assert "ENCABEZADO DE LA NOTARIA" in texto_completo
    assert "Cuerpo de la transcripcion." in texto_completo


def test_numeracion_de_paginas_opcional(tmp_path: Path):
    config = ConfiguracionApp(numeracion_paginas=True)
    paginas = [docx_export.PaginaTranscrita(numero=1, texto="Texto.")]
    ruta_salida = tmp_path / "salida.docx"
    docx_export.generar_documento(paginas, config, ruta_salida)

    documento = docx.Document(str(ruta_salida))
    xml = documento.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in xml
